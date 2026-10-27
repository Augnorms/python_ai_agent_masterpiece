import chromadb
from sentence_transformers import SentenceTransformer
import io
import docx
import pypdf

# Connect to Docker ChromaDB
chroma_client = chromadb.HttpClient(host='localhost', port=8001)
embedding_model = SentenceTransformer('all-MiniLM-L6-v2')

def extract_text_from_file(file_content: bytes, filename: str) -> str:
    """Extracts text based on file format (.docx, .pdf, .txt, .md, etc.)."""
    lower_name = filename.lower()
    
    # 1. Handle DOCX files
    if lower_name.endswith('.docx'):
        doc = docx.Document(io.BytesIO(file_content))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # Also include any table content
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs)
        
    # 2. Handle PDF files
    elif lower_name.endswith('.pdf'):
        reader = pypdf.PdfReader(io.BytesIO(file_content))
        pages_text = [page.extract_text().strip() for page in reader.pages if page.extract_text()]
        return "\n\n".join(pages_text)
        
    # 3. Handle Plain Text, Markdown, JSON, etc.
    else:
        try:
            return file_content.decode('utf-8')
        except UnicodeDecodeError:
            # Fallback for Windows-1252 / Latin-1 encoded text files
            return file_content.decode('latin-1', errors='ignore')

def process_and_ingest_file(file_content: bytes, filename: str):
    """Reads a file, chunks it, embeds it, and saves it to ChromaDB."""
    try:
        # 1. Extract plain text from the file (supporting binary docx, pdf, txt, etc.)
        text = extract_text_from_file(file_content, filename)
        
        if not text or not text.strip():
            return "File was empty or contained no extractable text."
        
        # 2. Chunk the text (split by paragraphs)
        chunks = [chunk.strip() for chunk in text.split("\n\n") if chunk.strip()]
        
        # Fallback if text does not contain double newlines
        if len(chunks) <= 1:
            line_chunks = [line.strip() for line in text.split("\n") if len(line.strip()) > 30]
            if len(line_chunks) > 1:
                chunks = line_chunks
        
        if not chunks:
            chunks = [text.strip()]

        # 3. Create embeddings
        embeddings = embedding_model.encode(chunks).tolist()
        
        # 4. Save to ChromaDB
        collection = chroma_client.get_or_create_collection(name="company_docs")
        
        # Create safe unique IDs based on filename and chunk index
        safe_filename = filename.replace(" ", "_")
        ids = [f"{safe_filename}_chunk_{i}" for i in range(len(chunks))]
        metadatas = [{"source": filename} for _ in chunks]
        
        collection.add(
            documents=chunks,
            embeddings=embeddings,
            ids=ids,
            metadatas=metadatas
        )
        
        return f"Successfully ingested {len(chunks)} chunks into ChromaDB."
        
    except Exception as e:
        raise Exception(f"Ingestion failed: {str(e)}")